#!/usr/bin/env python3
"""Export all available variables for one Synoptic viewer station to an .xlsx.

The viewer's public token is limited to the trailing 7 days of history, so this
pulls every sensor variable the station reports over that window and writes an
Excel workbook with a data sheet plus a station-info sheet.

Stdlib for fetching (no third-party requests); openpyxl for the workbook.
"""

import json
import os
import re
from datetime import datetime, timedelta, timezone
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Font

VIEWER_ROOT = "https://viewer.synopticdata.com/"
API_BASE = "https://api.synopticdata.com/v2"

STID = "AMB5847"
UNITS = "metric"
OBTIMEZONE = "UTC"
LOOKBACK_DAYS = 7

# Short variable names as accepted by the `vars` parameter, in display order.
VARS = [
    "air_temp",
    "relative_humidity",
    "dew_point_temperature",
    "heat_index",
    "PM_25_concentration",
    "PM_10_concentration",
    "ozone_concentration",
    "NO2_concentration",
    "CO_concentration",
    "air_quality_index_usa",
    "air_quality_index_eu",
    "air_quality_index_raw",
]

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(HERE, f"{STID}_full.xlsx")

UA = {"User-Agent": "Mozilla/5.0"}


def http_text(url):
    req = Request(url, headers=UA)
    with urlopen(req, timeout=30) as r:
        return r.read().decode("utf-8")


def http_json(url, params=None):
    if params:
        url = f"{url}?{urlencode(params)}"
    req = Request(url, headers=UA)
    with urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def discover_token():
    html = http_text(VIEWER_ROOT)
    m = re.search(r'src="(/assets/index-[A-Za-z0-9_-]+\.js)"', html)
    if not m:
        raise RuntimeError("Could not find the JS bundle in the viewer HTML")
    js = http_text(VIEWER_ROOT.rstrip("/") + m.group(1))
    t = re.search(r'token\s*:\s*["\']([A-Za-z0-9_-]+)["\']', js)
    if not t:
        raise RuntimeError("Could not find the embedded token in the JS bundle")
    return t.group(1)


def fetch_timeseries(token, start_dt):
    start = start_dt.strftime("%Y%m%d%H%M")
    end = datetime.now(timezone.utc).strftime("%Y%m%d%H%M")
    data = http_json(f"{API_BASE}/stations/timeseries", {
        "stid": STID, "token": token, "vars": ",".join(VARS),
        "start": start, "end": end,
        "units": UNITS, "obtimezone": OBTIMEZONE,
    })
    return data["STATION"][0]


def heat_index_c(temp_c, rh):
    """NOAA heat index (Celsius). Returns temp_c below the 80F/40% threshold."""
    t = temp_c * 9.0 / 5.0 + 32.0
    if t < 80 or rh < 40:
        return temp_c
    hi = (-42.379 + 2.04901523 * t + 10.14333127 * rh
          - 0.22475541 * t * rh - 6.83783e-3 * t * t - 5.481717e-2 * rh * rh
          + 1.22874e-3 * t * t * rh + 8.5282e-4 * t * rh * rh
          - 1.99e-6 * t * t * rh * rh)
    return round((hi - 32.0) * 5.0 / 9.0, 2)


def fetch_metadata(token):
    data = http_json(f"{API_BASE}/stations/metadata", {
        "stid": STID, "token": token,
    })
    return data["STATION"][0]


def main():
    token = discover_token()
    start_dt = datetime.now(timezone.utc) - timedelta(days=LOOKBACK_DAYS)

    station = fetch_timeseries(token, start_dt)
    obs = station["OBSERVATIONS"]
    dates = obs["date_time"]

    # Map each requested variable to its actual OBSERVATIONS key
    # (derived vars use a `_set_1d` suffix, primary sensors use `_set_1`).
    col_map = {}
    for var in VARS:
        for key in obs:
            if key.startswith(var + "_set_"):
                col_map[var] = key
                break

    meta = fetch_metadata(token)

    wb = Workbook()

    # --- Data sheet ---
    ws = wb.active
    ws.title = "data"
    headers = ["date_time"] + VARS
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    air_temp_key = col_map.get("air_temp")
    rh_key = col_map.get("relative_humidity")
    hi_key = col_map.get("heat_index")

    for i, dt in enumerate(dates):
        row = [dt]
        for var in VARS:
            key = col_map.get(var)
            if key is None:
                row.append(None)
                continue
            v = obs[key][i]
            # heat_index is null throughout the history response; compute it.
            if var == "heat_index" and (v is None):
                t = obs[air_temp_key][i] if air_temp_key else None
                r = obs[rh_key][i] if rh_key else None
                v = heat_index_c(t, r) if (t is not None and r is not None) else None
            # air_quality_index_raw is a nested JSON object; flatten to text.
            if isinstance(v, (dict, list)):
                v = json.dumps(v)
            row.append(v)
        ws.append(row)

    ws.freeze_panes = "A2"

    # --- Station info sheet ---
    info = wb.create_sheet("station_info")
    for k, v in meta.items():
        if isinstance(v, (dict, list)):
            v = json.dumps(v)
        info.append([k, v])
    info["A1"].font = Font(bold=True)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 60

    wb.save(OUT_XLSX)
    print(f"OK wrote={OUT_XLSX} rows={len(dates)} cols={len(VARS)} "
          f"window={dates[0]}..{dates[-1]}")


if __name__ == "__main__":
    main()
