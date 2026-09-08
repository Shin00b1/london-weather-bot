"""Backfill the unified trade tape into london/tape.duckdb (table `trades`).

Two sources, both deduplicated by INSERT OR IGNORE so the script is safe to
re-run and overlaps are harmless:

  1. Existing per-run snapshots london/runs/trades_*.jsonl (data-api schema).
  2. The data-api /trades endpoint, paged to FULL history for every market in
     the workbook with event_date >= CUTOFF. This closes the gap between the
     frozen legacy parquet (~2026-07-20) and the start of the run snapshots.

The legacy chain-level tape (london_temperature_unified_trades.parquet /
.csv.gz, 2.76M trades, ends ~2026-07-20) is left untouched — its schema
(maker/taker, block_number) differs from the data-api taker-side schema that
keeps growing nightly. CUTOFF defaults to the day after that legacy end.

Usage: python3 backfill_trades.py [--cutoff YYYY-MM-DD]
"""

import glob
import json
import os
import sys
import time
import urllib.error
import urllib.request
from datetime import datetime, timezone

import openpyxl

import tape

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "london", "london_temperature_unified.xlsx")
RUNS = os.path.join(BASE, "london", "runs")
DATA_API = "https://data-api.polymarket.com"
UA = {"User-Agent": "london-weather-updater/1.0"}

DEFAULT_CUTOFF = "2026-07-21"


def http_json(url, retries=3):
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=30) as r:
                return json.loads(r.read())
        except urllib.error.HTTPError as e:
            if e.code == 404:
                return None
            if attempt == retries - 1:
                raise
            time.sleep(2 * (attempt + 1))
        except Exception:
            if attempt == retries - 1:
                raise
            time.sleep(2 * (attempt + 1))


def markets_from_sheet(cutoff):
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb["Markets"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        ed = r[col["event_date"]]
        if ed and hasattr(ed, "date"):
            ed = ed.date()
        s = str(ed)[:10] if ed else ""
        if not s or s < cutoff:
            continue
        cond = r[col["condition_id"]]
        if not cond:
            continue
        out.append((cond.lower(), r[col["slug"]]))
    wb.close()
    # dedupe by condition_id preserving order
    seen = set()
    uniq = []
    for cond, slug in out:
        if cond not in seen:
            seen.add(cond)
            uniq.append((cond, slug))
    return uniq


def fetch_market_full(cond):
    """Page the data-api /trades endpoint to full history, newest-first."""
    rows, offset = [], 0
    while True:
        page = http_json(DATA_API + "/trades?market=%s&limit=500&offset=%d" % (cond, offset))
        if not page:
            break
        rows.extend(page)
        if len(page) < 500:
            break
        offset += 500
        time.sleep(0.25)
    return rows


def merge_run_files(con, fetched_at):
    total = 0
    for path in sorted(glob.glob(os.path.join(RUNS, "trades_*.jsonl"))):
        n = 0
        with open(path) as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    t = json.loads(line)
                except json.JSONDecodeError:
                    continue
                n += 1
                tape.append_trades(con, [t], fetched_at, src=os.path.basename(path))
        print("  run file %s: %d rows" % (os.path.basename(path), n))
        total += n
    return total


def main():
    cutoff = DEFAULT_CUTOFF
    if "--cutoff" in sys.argv:
        cutoff = sys.argv[sys.argv.index("--cutoff") + 1]
    fetched_at = datetime.now(timezone.utc)

    markets = markets_from_sheet(cutoff)
    print("markets with event_date >= %s: %d" % (cutoff, len(markets)))

    con = tape.connect()

    print("merging existing run files...")
    run_rows = merge_run_files(con, fetched_at)

    print("backfilling from data-api...")
    api_rows = 0
    for i, (cond, slug) in enumerate(markets, 1):
        try:
            trades = fetch_market_full(cond)
        except Exception as e:
            print("  FAIL %s: %r" % (slug, e))
            continue
        tape.append_trades(con, trades, fetched_at, src="data-api")
        api_rows += len(trades)
        if i % 50 == 0:
            print("  ...%d/%d markets, %d trades so far" % (i, len(markets), api_rows))
        time.sleep(0.15)

    n = con.execute("SELECT count(*) FROM trades").fetchone()[0]
    rng = con.execute("SELECT min(timestamp), max(timestamp) FROM trades").fetchone()
    con.close()
    print("DONE run_rows=%d api_rows=%d | table total=%d | ts range %s -> %s"
          % (run_rows, api_rows, n,
             datetime.utcfromtimestamp(rng[0]).isoformat() if rng[0] else None,
             datetime.utcfromtimestamp(rng[1]).isoformat() if rng[1] else None))


if __name__ == "__main__":
    main()
