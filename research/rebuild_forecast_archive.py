"""Rebuild the London temperature forecast archive with honest D-1 data.

Source: Open-Meteo Single Runs API, model `ecmwf_ifs` (ECMWF IFS HRES 9 km),
which archives each operational run back to 2024-03-14. For each event date we
fetch the run initialized at D-1 00:00 UTC and read the max/min forecast for
the target date — the value a trader could actually have known at a D-1 noon
entry (the 00Z run is the freshest deterministic run fully available then).

This replaces the look-ahead-contaminated `fcst_model_c` (historical-forecast
API, which stitched near-zero-lead nowcasts).

Outputs:
  london/forecast_archive_d1.csv   date,d1_ecmwf_max,d1_ecmwf_min
  london/forecast_archive_cache.json  (checkpoint, resumable)
"""

import json
import os
import sys
import time
import urllib.request
from datetime import date, datetime, timedelta

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
LONDON = os.path.join(BASE, "london")
XLSX = os.path.join(LONDON, "london_temperature_unified.xlsx")
CSV = os.path.join(LONDON, "forecast_archive_d1.csv")
CACHE = os.path.join(LONDON, "forecast_archive_cache.json")

LAT, LON = "51.5053", "0.0553"
MODEL = "ecmwf_ifs"


def get(url, retries=4, timeout=90):
    for a in range(retries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "london-rebuild/1.0"})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return json.loads(r.read())
        except Exception as e:
            if a == retries - 1:
                raise
            time.sleep(3 * (a + 1))


def d1_forecast(target):
    """D-1 00:00 UTC ecmwf_ifs forecast for `target`, returns (max, min)."""
    run = (target - timedelta(days=1)).isoformat()
    url = ("https://single-runs-api.open-meteo.com/v1/forecast"
           "?latitude=%s&longitude=%s"
           "&daily=temperature_2m_max,temperature_2m_min"
           "&run=%sT00:00&models=%s&timezone=UTC" % (LAT, LON, run, MODEL))
    d = get(url)
    if d.get("error"):
        raise RuntimeError(d.get("reason", "unknown"))
    t = d["daily"]["time"]
    mx = d["daily"]["temperature_2m_max"]
    mn = d["daily"]["temperature_2m_min"]
    if target.isoformat() not in t:
        raise RuntimeError("target not in horizon")
    i = t.index(target.isoformat())
    return mx[i], mn[i]


def event_dates():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb["Events"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    dates = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        ed = r[col["event_date"]]
        if isinstance(ed, datetime):
            ed = ed.date()
        if ed is not None:
            dates.add(ed)
    return sorted(dates)


def main():
    dates = event_dates()
    cache = {}
    if os.path.exists(CACHE):
        cache = json.load(open(CACHE))

    todo = [d for d in dates if d.isoformat() not in cache]
    print("total event-days: %d, already cached: %d, to fetch: %d"
          % (len(dates), len(dates) - len(todo), len(todo)))

    ok = fail = 0
    for i, d in enumerate(todo):
        ts = d.isoformat()
        try:
            mx, mn = d1_forecast(d)
            cache[ts] = {"max": mx, "min": mn}
            ok += 1
        except Exception as e:
            cache[ts] = None
            fail += 1
            print("  FAIL %s: %r" % (ts, e), file=sys.stderr)
        if (i + 1) % 25 == 0:
            json.dump(cache, open(CACHE, "w"), indent=1)
            print("  progress %d/%d (ok %d fail %d)" % (i + 1, len(todo), ok, fail))
        time.sleep(0.25)

    json.dump(cache, open(CACHE, "w"), indent=1)

    with open(CSV, "w") as f:
        f.write("date,d1_ecmwf_max,d1_ecmwf_min\n")
        for d in dates:
            ts = d.isoformat()
            c = cache.get(ts)
            if c:
                f.write("%s,%s,%s\n" % (ts, c["max"], c["min"]))
            else:
                f.write("%s,,\n" % ts)

    print("done. ok %d, fail %d -> %s" % (ok, fail, CSV))


if __name__ == "__main__":
    main()
