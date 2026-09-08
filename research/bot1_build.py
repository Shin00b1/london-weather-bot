#!/usr/bin/env python3
"""Bot 1 — build the bot1.xlsx workbook.

Consolidates the regime classifier, wind archive, and wind-veto analysis into
a single workbook at london/bot1.xlsx with sheets:

  README   bot spec: edge, regime split, vetos, validation status
  Regime   forecast-range vs lock-by-06Z sweep (the regime classifier)
  Daily    per-day joined record (forecast, actual, lock, wind, rain)
  Vetos    wind-veto candidates (direction shift, falling speed)
  Misses   known bot losses with wind/rain/regime context

Regime label uses the clean split: forecast range >= 8C = radiation (trade),
< 8C = advection (stand down). Wind direction is only meaningful at >= 6 kt.
"""
import csv
import math
import os
from collections import defaultdict
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
LONDON = os.path.join(HERE, "london")
OUT = os.path.join(LONDON, "bot1.xlsx")

REGIME_CSV = os.path.join(LONDON, "bot1_regime_daily.csv")
WIND_CSV = os.path.join(LONDON, "metar_eglc_wind.csv")
FCST_CSV = os.path.join(LONDON, "forecast_archive_d1.csv")
LOWBAND_CSV = os.path.join(LONDON, "lowband_archive.csv")

LON = ZoneInfo("Europe/London")
RANGE_THRESH = 8.0
DIR_MIN_WIND = 6.0  # direction meaningful only at >= this many knots


def load_fcst():
    out = {}
    with open(FCST_CSV) as f:
        for r in csv.DictReader(f):
            try:
                out[r["date"]] = (float(r["d1_ecmwf_max"]), float(r["d1_ecmwf_min"]))
            except (KeyError, ValueError):
                continue
    return out


def load_wind():
    days = defaultdict(list)
    with open(WIND_CSV) as f:
        for r in csv.DictReader(f):
            try:
                dt = datetime.strptime(r["valid"], "%Y-%m-%d %H:%M").replace(tzinfo=timezone.utc)
                sk = float(r["sknt"]) if r.get("sknt") else None
                dr = float(r["drct"]) if r.get("drct") else None
                days[dt.astimezone(LON).date()].append((dt, sk, dr))
            except (ValueError, KeyError):
                continue
    for d in days:
        days[d].sort()
    return days


def at(obs, hour):
    tgt = obs[0][0].replace(hour=hour, minute=0, second=0, microsecond=0)
    return min(obs, key=lambda o: abs((o[0] - tgt).total_seconds()))


def angdiff(a, b):
    if a is None or b is None:
        return None
    return (b - a + 180) % 360 - 180


def load_lowband():
    out = {}
    with open(LOWBAND_CSV) as f:
        for r in csv.DictReader(f):
            out[r["date"]] = r
    return out


def load_regime():
    rows = {}
    order = []
    with open(REGIME_CSV) as f:
        for r in csv.DictReader(f):
            rows[r["date"]] = r
            order.append(r["date"])
    return rows, order


def num(v):
    try:
        return float(v) if v not in ("", None) else None
    except (TypeError, ValueError):
        return None


def main():
    fcst = load_fcst()
    wind = load_wind()
    lowband = load_lowband()
    regime, order = load_regime()

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    # ---- README ----
    ws = wb.active
    ws.title = "README"
    readme = [
        ("BOT 1 — London lowest-temperature (dawn-lock) regime bot", ""),
        ("", ""),
        ("Edge", "Daily MIN is locked by 06:00 UTC on radiation-dominated days; "
                 "the market reprices that fact hours late. Trade the lag, not the forecast."),
        ("SIGNAL (corrected)", "The settlement IS the EGLC METAR, so the EGLC METAR's own pre-06Z "
                 "minimum is the answer's precursor — free, and 9/9 on recent settled days. The "
                 "private-station blend is now DROPPED from the call: it went 6/9 on the same days, "
                 "missing exactly the 3 the bot lost (Aug 22/24/26). Stations add correction noise, "
                 "not lead, because the market already lags the METAR by hours."),
        ("Failure mode", "Post-dawn cold front drops temp after 06Z and breaks the bucket."),
        ("Regime split", "Forecast diurnal range (D-1 ECMWF max-min) >= 8C = radiation (trade); "
                         "< 8C = advection (stand down). NOT a calendar split."),
        ("Entry", "05:30-07:00 London; read EGLC METAR min over the London day up to 06Z; "
                  "buy that bucket if the book still lags."),
        ("Veto 1 (band)", "Effective price 0.63-0.97 only."),
        ("Veto 2 (rain)", "Skip wet dawns (rain before 06Z): hit rate 71% vs 97% dry."),
        ("Veto 3 (wind)", "Skip if pre-dawn wind direction shifts >=10 deg (frontal veer/back) at >=6 kt, "
                          "or wind falls >=1 kt 03->06. CANDIDATE — needs forward validation."),
        ("Upgrade path", "50-member ECMWF ensemble spread replaces the single-point threshold; tight "
                         "spread = radiation regime by construction."),
        ("Status", "Regime classifier VALIDATED on 583 settled days. Signal corrected to METAR-only "
                   "(9/9 vs stations 6/9). Wind veto EXPLORATORY. Winter min tape does not exist; "
                   "classifier's winter job is stand-down."),
        ("Sheets", "Regime / Daily / Vetos / Misses"),
    ]
    for r in readme:
        ws.append(list(r))

    # ---- Regime ----
    ws = wb.create_sheet("Regime")
    ws.append(["forecast_range_threshold", "radiation_n", "radiation_lock_pct",
               "advection_n", "advection_lock_pct", "note"])
    scored = [regime[d] for d in order if num(regime[d]["fcst_range"]) is not None]
    for thr in [3, 4, 5, 6, 7, 8, 9, 10]:
        hi = [r for r in scored if num(r["fcst_range"]) >= thr]
        lo = [r for r in scored if num(r["fcst_range"]) < thr]
        if not hi or not lo:
            continue
        ws.append([thr, len(hi),
                   round(sum(int(r["lock06"]) for r in hi) / len(hi) * 100, 1),
                   len(lo),
                   round(sum(int(r["lock06"]) for r in lo) / len(lo) * 100, 1),
                   "TRADE / STAND DOWN" if thr == RANGE_THRESH else ""])
    ws.cell(1, 1).font = bold
    ws.append([])
    ws.append(["Range bin", "n", "lock%", "mean_actual_range", "mean_break_c"])
    bins = [(-99, 3), (3, 4), (4, 5), (5, 6), (6, 7), (7, 8), (8, 9), (9, 10), (10, 12), (12, 99)]
    for lo, hi in bins:
        g = [r for r in scored if lo <= num(r["fcst_range"]) < hi]
        if not g:
            continue
        ws.append([f"{lo}-{hi}", len(g),
                   round(sum(int(r["lock06"]) for r in g) / len(g) * 100, 1),
                   round(sum(num(r["actual_range"]) for r in g) / len(g), 1),
                   round(sum(num(r["break_c"]) for r in g) / len(g), 2)])

    # ---- Daily ----
    ws = wb.create_sheet("Daily")
    fields = ["date", "fcst_max", "fcst_min", "fcst_range", "actual_range",
              "pre06_min", "day_min", "break_c", "lock06", "regime",
              "sknt_06z", "ds_03", "dd_03", "dir_ok", "rain_pre06", "rain_2h",
              "call_bucket", "settle_min_c", "call_hit"]
    ws.append(fields)
    for i in range(1, len(fields) + 1):
        ws.cell(1, i).font = bold

    for d in order:
        r = regime[d]
        w = wind.get(datetime.strptime(d, "%Y-%m-%d").date(), [])
        fr = num(r["fcst_range"])
        regime_label = ("radiation" if fr is not None and fr >= RANGE_THRESH
                        else ("advection" if fr is not None else ""))
        sk6 = ds03 = dd03 = None
        dir_ok = 0
        if w:
            o3 = at(w, 3); o6 = at(w, 6)
            s3, dr3 = o3[1], o3[2]; s6, dr6 = o6[1], o6[2]
            sk6 = s6
            if s6 is not None and s3 is not None:
                ds03 = round(s6 - s3, 1)
            if dr3 is not None and dr6 is not None and s6 is not None and s6 >= DIR_MIN_WIND:
                dd03 = round(angdiff(dr3, dr6), 1)
                dir_ok = 1
        lb = lowband.get(d, {})
        ws.append([d,
                   fcst.get(d, ("", ""))[0], fcst.get(d, ("", ""))[1],
                   fr, num(r["actual_range"]), num(r["pre06_min"]), num(r["day_min"]),
                   num(r["break_c"]), int(r["lock06"]), regime_label,
                   sk6, ds03, dd03, dir_ok,
                   lb.get("rain_pre06", ""), lb.get("rain_2h", ""),
                   lb.get("call_bucket", ""), lb.get("settle_min_c", ""),
                   lb.get("call_hit", "")])

    # ---- Vetos ----
    ws = wb.create_sheet("Vetos")
    ws.append(["veto", "threshold", "breaks_skipped", "breaks_total", "breaks_pct",
               "locks_cost", "locks_total", "locks_pct"])
    ws.cell(1, 1).font = bold
    brk = [r for r in scored if int(r["lock06"]) == 0]
    lok = [r for r in scored if int(r["lock06"]) == 1]
    # build wind features per date for the veto sweep
    feat = {}
    for d in order:
        w = wind.get(datetime.strptime(d, "%Y-%m-%d").date(), [])
        if not w:
            continue
        o3 = at(w, 3); o6 = at(w, 6)
        s3, dr3 = o3[1], o3[2]; s6, dr6 = o6[1], o6[2]
        feat[d] = {"sk6": s6,
                   "ds03": (s6 - s3) if s6 is not None and s3 is not None else None,
                   "dd03": angdiff(dr3, dr6) if dr3 is not None and dr6 is not None else None,
                   "dir_ok": s6 is not None and s6 >= DIR_MIN_WIND and dr3 is not None and dr6 is not None}

    def veto_row(name, thresh, cond):
        bb = [d for d in feat if d in (x["date"] for x in brk) and cond(feat[d]) is not None]
        # restrict breaks/locks to days with a computable feature
        brk_dates = {r["date"] for r in brk}
        lok_dates = {r["date"] for r in lok}
        bb_tot = [d for d in feat if d in brk_dates and cond(feat[d]) is not None]
        ll_tot = [d for d in feat if d in lok_dates and cond(feat[d]) is not None]
        sb = sum(1 for d in bb_tot if cond(feat[d]))
        sl = sum(1 for d in ll_tot if cond(feat[d]))
        ws.append([name, thresh, sb, len(bb_tot),
                   round(sb / len(bb_tot) * 100, 1) if bb_tot else "",
                   sl, len(ll_tot),
                   round(sl / len(ll_tot) * 100, 1) if ll_tot else ""])

    for x in [10, 20, 30]:
        veto_row("|dir shift| 03->06 >= deg", x, lambda f, x=x: (abs(f["dd03"]) >= x) if f["dir_ok"] else None)
    for x in [1, 2, 3]:
        veto_row("speed fall 03->06 <= -kt", x, lambda f, x=x: (f["ds03"] <= -x) if f["ds03"] is not None else None)

    # ---- Misses ----
    ws = wb.create_sheet("Misses")
    ws.append(["date", "station_call", "metar_pre06_bucket", "settle_min_c",
               "station_wrong", "metar_right", "break_c", "rain_pre06", "rain_2h",
               "sknt_06z", "ds_03", "dd_03", "regime", "wind_flag", "rain_flag"])
    ws.cell(1, 1).font = bold
    miss_dates = [d for d in order
                  if lowband.get(d, {}).get("call_hit") == "0"]
    for d in miss_dates:
        lb = lowband[d]
        r = regime[d]
        pre06 = num(r["pre06_min"])
        settle = num(lb.get("settle_min_c"))
        metar_bucket = round(pre06) if pre06 is not None else ""
        settle_bucket = round(settle) if settle is not None else ""
        w = wind.get(datetime.strptime(d, "%Y-%m-%d").date(), [])
        sk6 = ds03 = dd03 = None
        if w:
            o3 = at(w, 3); o6 = at(w, 6)
            s3, dr3 = o3[1], o3[2]; s6, dr6 = o6[1], o6[2]
            sk6 = s6
            if s6 is not None and s3 is not None:
                ds03 = round(s6 - s3, 1)
            if dr3 is not None and dr6 is not None and s6 is not None and s6 >= DIR_MIN_WIND:
                dd03 = round(angdiff(dr3, dr6), 1)
        fr = num(r["fcst_range"])
        regime_label = ("radiation" if fr is not None and fr >= RANGE_THRESH
                        else ("advection" if fr is not None else ""))
        wind_flag = ("DIR" if (dd03 is not None and abs(dd03) >= 10)
                     else ("FALL" if (ds03 is not None and ds03 <= -1) else ""))
        rain_flag = "RAIN" if lb.get("rain_pre06") == "1" else ""
        ws.append([d, lb.get("call_bucket", ""), metar_bucket, settle,
                   int(lb.get("call_bucket", "") != str(settle_bucket) if lb.get("call_bucket") else 0),
                   int(metar_bucket == settle_bucket and settle_bucket != ""),
                   num(r["break_c"]), lb.get("rain_pre06", ""), lb.get("rain_2h", ""),
                   sk6, ds03, dd03, regime_label, wind_flag, rain_flag])

    wb.save(OUT)
    print("OK wrote %s" % OUT)
    print("  sheets: %s" % ", ".join(wb.sheetnames))
    print("  regime rows=%d  daily rows=%d  misses=%d" % (len(scored), len(order), len(miss_dates)))


if __name__ == "__main__":
    main()
