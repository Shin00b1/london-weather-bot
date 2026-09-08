"""One-off cleanup for london_temperature_unified.xlsx and run artifacts.

- Markets: dedupe by condition_id (keep the most recently written row),
  sort by (event_date, direction, bucket_low_c)
- Events:  dedupe by event_slug, sort by (event_date, direction)
- Daily:   dedupe by trade_date, sort ascending
- runs/:   dedupe trade lines by (txHash, wallet, ts, asset) and gzip them

Only touches Markets/Events/Daily data rows; README/Reconciliation/
FreeLunch/FavoriteAnalysis are left untouched. Backs up the workbook first.
"""

import csv
import gzip
import json
import os
import re
import shutil
from datetime import datetime

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
LONDON = os.path.join(BASE, "london")
XLSX = os.path.join(LONDON, "london_temperature_unified.xlsx")
BACKUPS = os.path.join(LONDON, "backups")
RUNS = os.path.join(LONDON, "runs")

BUCKET_NONE = -999.0


def sheet_rows(ws):
    hdr = [c.value for c in ws[1]]
    return hdr, [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def rewrite(ws, hdr, rows):
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        ws.append(r)
    # keep header exact
    for i, h in enumerate(hdr, 1):
        ws.cell(1, i, h)


def as_date(v):
    return v.date() if isinstance(v, datetime) else v


def main():
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(XLSX, os.path.join(BACKUPS, "unified_preclean_%s.xlsx" % ts))
    wb = openpyxl.load_workbook(XLSX)

    # Markets
    ws = wb["Markets"]
    hdr, rows = sheet_rows(ws)
    i = {h: n for n, h in enumerate(hdr)}
    uniq = {}
    for r in rows:
        uniq[r[i["condition_id"]]] = r  # later rows win
    def mkey(r):
        low = r[i["bucket_low_c"]]
        return (as_date(r[i["event_date"]]) or date_min,
                r[i["direction"]] or "",
                float(low) if low is not None else BUCKET_NONE)
    date_min = min(as_date(r[i["event_date"]]) for r in uniq.values() if r[i["event_date"]])
    cleaned = sorted(uniq.values(), key=mkey)
    rewrite(ws, hdr, cleaned)
    print("Markets: %d -> %d rows" % (len(rows), len(cleaned)))

    # Events
    ws = wb["Events"]
    hdr, rows = sheet_rows(ws)
    i = {h: n for n, h in enumerate(hdr)}
    uniq = {}
    for r in rows:
        uniq[r[i["event_slug"]]] = r
    cleaned = sorted(uniq.values(),
                     key=lambda r: (as_date(r[i["event_date"]]) or date_min, r[i["direction"]] or ""))
    rewrite(ws, hdr, cleaned)
    print("Events: %d -> %d rows" % (len(rows), len(cleaned)))

    # Daily
    ws = wb["Daily"]
    hdr, rows = sheet_rows(ws)
    i = {h: n for n, h in enumerate(hdr)}
    uniq = {}
    for r in rows:
        uniq[as_date(r[i["trade_date"]])] = r
    cleaned = sorted(uniq.values(), key=lambda r: as_date(r[i["trade_date"]]))
    rewrite(ws, hdr, cleaned)
    print("Daily: %d -> %d rows" % (len(rows), len(cleaned)))

    tmp = XLSX + ".tmp"
    wb.save(tmp)
    os.replace(tmp, XLSX)
    print("workbook saved")

    # compress raw run trades
    for fn in sorted(os.listdir(RUNS)):
        if not fn.startswith("trades_") or not fn.endswith(".jsonl"):
            continue
        path = os.path.join(RUNS, fn)
        seen, kept = set(), 0
        with open(path) as f, gzip.open(path + ".gz", "wt") as out:
            for line in f:
                try:
                    d = json.loads(line)
                except Exception:
                    continue
                k = (d.get("transactionHash"), d.get("proxyWallet"),
                     d.get("timestamp"), d.get("asset"))
                if k in seen:
                    continue
                seen.add(k)
                out.write(line)
                kept += 1
        os.remove(path)
        print("%s: deduped+gzipped -> %d unique trades" % (fn, kept))


if __name__ == "__main__":
    main()
