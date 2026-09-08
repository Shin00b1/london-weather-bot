"""Compare strategy edge using archive (look-ahead) vs genuine D-1 forecast.

Strategy (matching the deck): for each resolved event, buy the single bucket
containing the model forecast (round to nearest integer, clamp to edge
buckets), filter to buckets priced 0.10-0.50 at open, hold to resolution.
PnL per $1 = (1-p)/p if won else -1.0.

Two forecast sources:
  - archive: fcst_model_c from the sheet (Open-Meteo historical best_match,
    near-zero-lead -> look-ahead contaminated)
  - d1: genuine forecast issued the prior day (Single Runs API, ukmo_seamless
    run = target-1 00:00 UTC)

Covers the 224 resolved events since 2026-05-01 where both forecasts exist.
"""

import json
import os
from datetime import datetime, date

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "london", "london_temperature_unified.xlsx")
RES = os.path.join(BASE, "london", "lookahead_results.json")


def load_markets():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb["Markets"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        ed = r[col["event_date"]]
        if isinstance(ed, datetime):
            ed = ed.date()
        rows.append({
            "event_date": ed,
            "direction": r[col["direction"]],
            "bucket_low": r[col["bucket_low_c"]],
            "bucket_high": r[col["bucket_high_c"]],
            "bucket_label": r[col["bucket_label"]],
            "open_prob": r[col["open_prob"]],
            "won": r[col["won"]],
        })
    return rows


def forecast_to_bucket(fc, direction, markets):
    """Round forecast to nearest int, find the bucket containing it."""
    r = round(fc)
    for m in markets:
        lo, hi = m["bucket_low"], m["bucket_high"]
        if lo is None and hi is None:
            continue
        # edge buckets: "18 or below" has lo==hi==18 but is a <=18 bound
        label = m["bucket_label"] or ""
        if direction == "highest":
            if "below" in label and r <= hi:
                return m
            if "higher" in label and r >= lo:
                return m
            if lo == r == hi:
                return m
        else:
            if "below" in label and r <= hi:
                return m
            if "higher" in label and r >= lo:
                return m
            if lo == r == hi:
                return m
    # fallback: exact single-degree bucket
    for m in markets:
        if m["bucket_low"] == r and m["bucket_high"] == r:
            return m
    return None


def run_strategy(forecasts):
    """forecasts: {(date, direction): fc_value}. Returns stats."""
    markets = load_markets()
    by_event = {}
    for m in markets:
        by_event.setdefault((m["event_date"], m["direction"]), []).append(m)

    trades = []
    for (ed, direction), fc in forecasts.items():
        ms = by_event.get((ed, direction), [])
        if not ms or fc is None:
            continue
        bucket = forecast_to_bucket(fc, direction, ms)
        if bucket is None:
            continue
        p = bucket["open_prob"]
        if p is None or p < 0.10 or p > 0.50:
            continue
        won = bucket["won"]
        pnl = (1 - p) / p if won else -1.0
        trades.append((p, won, pnl))

    n = len(trades)
    if n == 0:
        return {"n": 0}
    wins = sum(1 for t in trades if t[1])
    pnl_per_dollar = sum(t[2] for t in trades) / n
    return {
        "n": n,
        "win_rate": wins / n,
        "pnl_per_dollar": pnl_per_dollar,
        "total_pnl": sum(t[2] for t in trades),
    }


def main():
    # archive forecasts from the sheet
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb["Events"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    archive_fc = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        ed = r[col["event_date"]]
        if isinstance(ed, datetime):
            ed = ed.date()
        fc = r[col["fcst_model_c"]]
        direction = r[col["direction"]]
        if fc is not None:
            archive_fc[(ed, direction)] = float(fc)

    # d1 forecasts from lookahead results
    d1_fc = {}
    for row in json.load(open(RES)):
        ed = date.fromisoformat(row["date"])
        d1_fc[(ed, row["direction"])] = float(row["d1"])

    # restrict archive to same event set as d1 (both sources present)
    common = set(d1_fc) & set(archive_fc)
    arch_sub = {k: archive_fc[k] for k in common}
    d1_sub = {k: d1_fc[k] for k in common}

    print("events with both forecasts:", len(common))
    print()
    print("=== ARCHIVE (look-ahead) forecast strategy ===")
    print(run_strategy(arch_sub))
    print()
    print("=== GENUINE D-1 forecast strategy ===")
    print(run_strategy(d1_sub))


if __name__ == "__main__":
    main()
