"""Annotate the ModelCompare sheet as superseded by the look-ahead finding.

The ModelCompare sheet was built from openmeteo_models.csv, sourced from the
Open-Meteo historical-forecast API (best_match / ukmo_seamless / globals).
That API stitches near-zero-lead nowcasts, so its "model beats market, MAE
0.47" headline is look-ahead-contaminated.

This script prepends a clear superseded banner to the ModelCompare sheet and
leaves the rest of the sheet intact for provenance. Honest D-1 numbers live in
the Calibration sheet (Table 2) and london/forecast_archive_d1.csv.

Backs up the workbook first (same convention as the other scripts).
"""

import os
import shutil
from datetime import datetime

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
LONDON = os.path.join(BASE, "london")
XLSX = os.path.join(LONDON, "london_temperature_unified.xlsx")
BACKUPS = os.path.join(LONDON, "backups")

BANNER = [
    "SUPERSEDED — this sheet used the Open-Meteo historical-forecast archive, which is look-ahead biased",
    "(",
    "  The historical-forecast API stitches each run's first hours = near-zero-lead nowcasts,",
    "  so its 'model MAE 0.47 / model beats market' result used forecasts not knowable at D-1.",
    "  Honest D-1 reconstruction (ECMWF IFS HRES, Single Runs API) gives model MAE ~1.0C",
    "  and NO significant edge vs the market. See Calibration Table 2 and forecast_archive_d1.csv.",
    ")",
]


def main():
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(XLSX, os.path.join(BACKUPS, "unified_preModelCompareNote_%s.xlsx" % ts))

    wb = openpyxl.load_workbook(XLSX)
    if "ModelCompare" not in wb.sheetnames:
        print("ModelCompare sheet not found; nothing to annotate")
        return
    ws = wb["ModelCompare"]
    ws.insert_rows(1, amount=len(BANNER))
    for i, line in enumerate(BANNER, start=1):
        ws.cell(i, 1, line)

    tmp = XLSX + ".tmp"
    wb.save(tmp)
    os.replace(tmp, XLSX)
    print("ModelCompare annotated as superseded (%d banner rows)" % len(BANNER))


if __name__ == "__main__":
    main()
